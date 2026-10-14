import json
import os
from datetime import datetime
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.utils.cell import range_boundaries

from .utils import setup_logging

logger = setup_logging(__name__)

class SpreadsheetAgent:
    """Execute spreadsheet tasks in the order specified by a JSON config."""
    def __init__(self, config):
        self.config = config
        self.workbook = None
        self.audit_log = []

    def audit(self, status, task_type, message, details=None):
        entry = {
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "status": status,
            "task_type": task_type,
            "message": message,
            "details": details or {}
        }
        self.audit_log.append(entry)
        getattr(logger, "error" if status == "ERROR" else "info")(
            "%s | %s | %s", status, task_type, message
        )

    def sheet(self, name):
        if name not in self.workbook.sheetnames:
            self.workbook.create_sheet(name)
        return self.workbook[name]

    def task_copy_data(self, task):
        source = task.get("source_sheet")
        destination = task.get("dest_sheet")
        if not source or not destination:
            raise ValueError("copy_data ต้องระบุ source_sheet และ dest_sheet")
        if source not in self.workbook.sheetnames:
            raise ValueError(f"ไม่พบชีตต้นทาง: {source}")
        src = self.workbook[source]
        dst = self.sheet(destination)
        # Clear old values before copying, so reruns do not leave stale rows.
        if dst.max_row:
            dst.delete_rows(1, dst.max_row)
        for row in src.iter_rows():
            for cell in row:
                target = dst[cell.coordinate]
                target.value = cell.value
                if cell.has_style:
                    target._style = copy(cell._style)
                if cell.number_format:
                    target.number_format = cell.number_format
        dst.freeze_panes = "A2"
        dst.auto_filter.ref = dst.dimensions
        self.audit("SUCCESS", "copy_data", f"คัดลอก {source} ไปยัง {destination}")

    def task_calculate_column(self, task):
        name, col, formula = task.get("sheet"), task.get("target_column"), task.get("formula")
        if not all([name, col, formula]):
            raise ValueError("calculate_column ต้องมี sheet, target_column และ formula")
        ws = self.sheet(name)
        start = int(task.get("start_row", 2))
        ws[f"{col}1"] = task.get("header", "Calculated Column")
        ws[f"{col}1"].font = Font(bold=True, color="FFFFFF")
        ws[f"{col}1"].fill = PatternFill("solid", fgColor="244062")
        count = 0
        for r in range(start, ws.max_row + 1):
            ws[f"{col}{r}"] = formula.replace("{row}", str(r))
            if task.get("number_format"):
                ws[f"{col}{r}"].number_format = task["number_format"]
            count += 1
        self.audit("SUCCESS", "calculate_column", f"เพิ่มสูตรคอลัมน์ {col} จำนวน {count} แถว")

    def task_conditional_format(self, task):
        ws = self.sheet(task["sheet"])
        cell_range, rule_type = task["range"], task["rule_type"]
        if rule_type == "color_scale":
            rule = ColorScaleRule(
                start_type="min", start_color=task.get("min_color", "FFF8696B"),
                mid_type="percentile", mid_value=50, mid_color=task.get("mid_color", "FFFFEB84"),
                end_type="max", end_color=task.get("max_color", "FF63BE7B"))
        elif rule_type == "data_bar":
            rule = DataBarRule(start_type="min", end_type="max",
                               color=task.get("color", "FF5B9BD5"), showValue=True)
        elif rule_type == "expression":
            formula = task.get("formula")
            if not formula:
                raise ValueError("expression ต้องมี formula")
            rule = FormulaRule(formula=[formula],
                fill=PatternFill("solid", fgColor=task.get("fill_color", "FFFFC7CE")),
                font=Font(color=task.get("font_color", "FF9C0006")))
        else:
            raise ValueError(f"ไม่รองรับ rule_type: {rule_type}")
        ws.conditional_formatting.add(cell_range, rule)
        self.audit("SUCCESS", "conditional_format", f"ใช้ {rule_type} กับ {task['sheet']}!{cell_range}")

    def task_chart(self, task):
        ws = self.sheet(task["sheet"])
        chart_type = task.get("chart_type", "bar")
        chart_classes = {"bar": BarChart, "line": LineChart, "pie": PieChart}
        if chart_type not in chart_classes:
            raise ValueError(f"ไม่รองรับ chart_type: {chart_type}")
        min_col, min_row, max_col, max_row = range_boundaries(task["data_range"])
        chart = chart_classes[chart_type]()
        data = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
        chart.add_data(data, titles_from_data=bool(task.get("titles_from_data", True)))
        category_range = task.get("category_range")
        if category_range:
            cmin, rmin, cmax, rmax = range_boundaries(category_range)
            chart.set_categories(Reference(ws, min_col=cmin, min_row=rmin, max_col=cmax, max_row=rmax))
        chart.title = task.get("title", "Sales Chart")
        if chart_type != "pie":
            chart.x_axis.title = task.get("x_axis_title", "Category")
            chart.y_axis.title = task.get("y_axis_title", "Value")
        chart.height = float(task.get("height", 8))
        chart.width = float(task.get("width", 14))
        ws.add_chart(chart, task.get("top_left_cell", "G2"))
        self.audit("SUCCESS", "chart", f"สร้างกราฟ {chart_type}: {chart.title}")

    def task_summary(self, task):
        source_name = task.get("source_sheet", "Raw Data")
        summary_name = task.get("sheet", "Summary")
        product_col = task.get("product_column", "A")
        revenue_col = task.get("revenue_column", "D")
        if source_name not in self.workbook.sheetnames:
            raise ValueError(f"ไม่พบชีตสำหรับสรุป: {source_name}")
        src = self.workbook[source_name]
        totals = {}
        for row in range(2, src.max_row + 1):
            product = src[f"{product_col}{row}"].value
            if product is None:
                continue
            # Excel formula cells may not have cached results until opened in Excel.
            # Calculate revenue from quantity * unit price for this sample workflow.
            revenue = src[f"{revenue_col}{row}"].value
            if isinstance(revenue, (int, float)):
                totals[str(product)] = totals.get(str(product), 0) + revenue
            elif product_col == "A" and src.max_column >= 3:
                qty, price = src[f"B{row}"].value, src[f"C{row}"].value
                if isinstance(qty, (int, float)) and isinstance(price, (int, float)):
                    totals[str(product)] = totals.get(str(product), 0) + qty * price
        out = self.sheet(summary_name)
        if out.max_row:
            out.delete_rows(1, out.max_row)
        out.append(["Product", "Total Revenue"])
        for product, total in sorted(totals.items(), key=lambda item: item[1], reverse=True):
            out.append([product, round(total, 2)])
        for c in out[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="244062")
        for r in range(2, out.max_row + 1):
            out[f"B{r}"].number_format = '#,##0.00'
        out.freeze_panes = "A2"
        out.auto_filter.ref = out.dimensions
        self.audit("SUCCESS", "summary", f"สรุปรายได้สินค้า {len(totals)} รายการ")

    def add_audit_sheet(self):
        name = "Audit Log"
        if name in self.workbook.sheetnames:
            del self.workbook[name]
        ws = self.workbook.create_sheet(name, 0)
        ws.append(["Timestamp", "Status", "Task Type", "Message", "Details"])
        for item in self.audit_log:
            ws.append([item["timestamp"], item["status"], item["task_type"],
                       item["message"], json.dumps(item["details"], ensure_ascii=False)])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="333333")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            letter = col[0].column_letter
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 55)

    def run(self):
        input_path = self.config["input_file"]
        output_path = self.config["output_file"]
        try:
            if not os.path.isfile(input_path):
                raise FileNotFoundError(f"ไม่พบไฟล์ Excel ต้นทาง: {input_path}")
            self.workbook = load_workbook(input_path)
            self.audit("SUCCESS", "load_workbook", f"เปิดไฟล์ {input_path}")
            dispatch = {
                "copy_data": self.task_copy_data,
                "calculate_column": self.task_calculate_column,
                "conditional_format": self.task_conditional_format,
                "chart": self.task_chart,
                "summary": self.task_summary,
            }
            for number, task in enumerate(self.config["tasks"], start=1):
                task_type = task["type"]
                try:
                    if task_type not in dispatch:
                        raise ValueError(f"ไม่รองรับ task type: {task_type}")
                    dispatch[task_type](task)
                except Exception as exc:
                    self.audit("ERROR", task_type, f"Task {number} ล้มเหลว: {exc}", task)
            self.add_audit_sheet()
            os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
            self.workbook.save(output_path)
            logger.info("บันทึกรายงานแล้ว: %s", output_path)
            return output_path
        except Exception as exc:
            self.audit("CRITICAL", "run", str(exc))
            raise
