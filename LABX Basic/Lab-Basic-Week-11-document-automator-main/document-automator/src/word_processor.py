# document-automator/src/word_processor.py

import os
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from utils import setup_logging, ensure_directory_exists

logger = setup_logging(__name__)


class WordProcessor:
    """Handles operations related to Word documents using python-docx."""

    def __init__(self):
        logger.info("WordProcessor initialized.")

    def create_document(self):
        """Creates a new blank Word document."""
        document = docx.Document()
        logger.info("New Word document created.")
        return document

    def add_heading(self, document, text, level=1):
        """Adds a heading to the document."""
        if not 0 <= level <= 9:
            logger.error("Heading level must be between 0 and 9.")
            return False

        document.add_heading(str(text), level=level)
        logger.info(f"Added heading (Level {level}): '{text}'")
        return True

    def add_paragraph(
        self,
        document,
        text,
        style=None,
        bold=False,
        italic=False,
        font_size=None,
        align="left"
    ):
        """Adds a paragraph with optional formatting."""
        paragraph = document.add_paragraph(style=style)
        run = paragraph.add_run(str(text))

        run.bold = bold
        run.italic = italic

        if font_size is not None:
            if font_size <= 0:
                logger.warning("Font size must be positive.")
            else:
                run.font.size = Pt(font_size)

        alignments = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }

        paragraph.alignment = alignments.get(
            align, WD_ALIGN_PARAGRAPH.LEFT
        )

        logger.info(f"Added paragraph: '{str(text)[:50]}...'")
        return paragraph

    def add_table(
        self,
        document,
        data,
        include_header=True,
        header_bold=True,
        col_widths=None
    ):
        """Adds a formatted table to the document."""

        if not data or not isinstance(data, (list, tuple)):
            logger.warning("No valid data provided to add table.")
            return None

        cols = len(data[0]) if isinstance(data[0], (list, tuple)) else 0

        if cols == 0:
            logger.warning("Table data is empty or malformed.")
            return None

        if any(
            not isinstance(row, (list, tuple)) or len(row) != cols
            for row in data
        ):
            logger.error("All table rows must have the same number of columns.")
            return None

        table = document.add_table(rows=len(data), cols=cols)
        table.style = "Table Grid"

        if col_widths is not None:
            if len(col_widths) != cols:
                logger.warning("Column widths do not match the table columns.")
            else:
                for index, width in enumerate(col_widths):
                    if width > 0:
                        table.columns[index].width = Inches(width)

        for row_index, row_data in enumerate(data):
            for col_index, cell_data in enumerate(row_data):
                cell = table.cell(row_index, col_index)
                cell.text = str(cell_data)

                if include_header and row_index == 0 and header_bold:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True

        logger.info(
            f"Added table with {len(data)} rows and {cols} columns."
        )
        return table

    def add_picture(self, document, image_path, width_inches=None):
        """Adds an image to the document with optional width scaling."""

        if not os.path.exists(image_path):
            logger.error(f"Image file not found: {image_path}")
            return False

        try:
            if width_inches is not None:
                if width_inches <= 0:
                    logger.error("Image width must be positive.")
                    return False

                document.add_picture(
                    image_path, width=Inches(width_inches)
                )
            else:
                document.add_picture(image_path)

            logger.info(f"Added picture: {os.path.basename(image_path)}")
            return True

        except Exception as e:
            logger.error(f"Error adding picture '{image_path}': {e}")
            return False

    def save_document(self, document, output_path):
        """Saves the Word document to the specified path."""

        try:
            output_dir = os.path.dirname(output_path)

            if output_dir:
                ensure_directory_exists(output_dir)

            document.save(output_path)

            logger.info(f"Word document saved to '{output_path}'.")
            return True

        except Exception as e:
            logger.error(
                f"Error saving Word document to '{output_path}': {e}"
            )
            return False